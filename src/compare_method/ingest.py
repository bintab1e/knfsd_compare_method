from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import yaml
from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from .common import (
    canonical_json,
    json_scalar,
    ordered_unique,
    repo_root,
    sha256_bytes,
    sha256_file,
    write_json,
)

SCHEMA_VERSION = "1.0.0"
HEADER_ALIASES = {
    "No.": "source_number",
    "이름": "title",
    "대상 구현체": "target_implementation",
    "NFS 버전": "protocol_version",
    "프로토콜·버전": "protocol_version",
    "우선 순위": "priority_raw",
    "기능/공격 표면": "attack_surface",
    "Root Cause": "root_cause",
    "영향 버전": "affected_versions",
    "CVSS": "cvss_raw",
    "참조 링크": "reference_url",
    "원문/패치 링크": "patch_related_url",
    "패치/관련 링크": "patch_related_url",
    "공개/커밋 날짜": "published_or_commit_date",
}
PRIORITY_MAP = {
    "매우 높음": "very_high",
    "높음": "high",
    "조건부 높음": "conditional_high",
    "보통": "medium",
    "낮음": "low",
    "검토 필요": "review_required",
    "미부여": "unassigned",
}
CVE_RE = re.compile(r"\bCVE-\d{4}-\d{4,7}\b", re.IGNORECASE)
COMMIT_URL_RE = re.compile(
    r"/(?:commit|commits?|linus|c)/(?:\?id=)?([0-9a-f]{7,64})(?:\b|/)", re.IGNORECASE
)
BARE_COMMIT_RE = re.compile(r"(?<![0-9a-f])([0-9a-f]{40})(?![0-9a-f])", re.IGNORECASE)


def _args() -> argparse.Namespace:
    root = repo_root()
    parser = argparse.ArgumentParser(description="Import the eight NFS 1-day workbooks.")
    parser.add_argument("--repo", type=Path, default=root)
    parser.add_argument("--input-dir", type=Path)
    parser.add_argument("--output-dir", type=Path)
    return parser.parse_args()


def _identifiers(values: list[Any]) -> tuple[list[str], list[str]]:
    text = "\n".join(str(v) for v in values if v not in (None, ""))
    cves = sorted({m.group(0).upper() for m in CVE_RE.finditer(text)})
    commits = {m.group(1).lower() for m in COMMIT_URL_RE.finditer(text)}
    commits.update(m.group(1).lower() for m in BARE_COMMIT_RE.finditer(text))
    return cves, sorted(commits)


def _cell_evidence(cell: Any) -> dict[str, Any]:
    return {
        "coordinate": cell.coordinate,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
        "style_id": cell.style_id,
        "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
        "comment": (
            {"author": cell.comment.author, "text": cell.comment.text}
            if cell.comment
            else None
        ),
    }


def _record(
    path: Path,
    workbook_sha256: str,
    implementation: str,
    sheet: Any,
    row_index: int,
    headers: list[str],
) -> dict[str, Any]:
    cells = [sheet.cell(row=row_index, column=i + 1) for i in range(len(headers))]
    raw_fields = {header: json_scalar(cell.value) for header, cell in zip(headers, cells)}
    fields: dict[str, Any] = {}
    for header, value in raw_fields.items():
        fields[HEADER_ALIASES[header]] = value

    source_number = fields.get("source_number")
    safe_number = re.sub(
        r"[^A-Za-z0-9._-]+",
        "-",
        str(source_number).strip() if source_number not in (None, "") else f"row-{row_index}",
    )
    record_id = f"{implementation}:{sheet.title}:{safe_number}:R{row_index}"
    source_cells = {header: _cell_evidence(cell) for header, cell in zip(headers, cells)}
    values = list(raw_fields.values()) + [
        cell.hyperlink.target for cell in cells if cell.hyperlink and cell.hyperlink.target
    ]
    cves, commits = _identifiers(values)

    references = []
    for kind, canonical in (
        ("reference", "reference_url"),
        ("patch_or_related", "patch_related_url"),
    ):
        header = next((h for h in headers if HEADER_ALIASES[h] == canonical), None)
        if header is None:
            continue
        cell = cells[headers.index(header)]
        display = json_scalar(cell.value)
        url = cell.hyperlink.target if cell.hyperlink else display
        if display not in (None, "") or url not in (None, ""):
            references.append(
                {
                    "kind": kind,
                    "display": None if display in (None, "") else str(display),
                    "url": None if url in (None, "") else str(url),
                    "source_record_id": record_id,
                    "source_cell": cell.coordinate,
                }
            )

    return {
        "record_id": record_id,
        "workbook": path.name,
        "workbook_sha256": workbook_sha256,
        "sheet": sheet.title,
        "excel_row": row_index,
        "source_number": source_number,
        "raw_fields": raw_fields,
        "source_cells": source_cells,
        "_normalized": fields,
        "_cves": cves,
        "_commits": commits,
        "_references": references,
    }


def _group_key(record: dict[str, Any]) -> tuple[str, str]:
    if record["_cves"]:
        return "cve", record["_cves"][0]
    if record["_commits"]:
        return "commit", record["_commits"][0]
    unresolved_payload = {
        "workbook": record["workbook"],
        "sheet": record["sheet"],
        "row": record["excel_row"],
        "raw_fields": record["raw_fields"],
    }
    return "unresolved", sha256_bytes(canonical_json(unresolved_payload))[:12].upper()


def _values(records: list[dict[str, Any]], field: str) -> list[str]:
    return ordered_unique([record["_normalized"].get(field) for record in records])


def _build_case(implementation: str, key: tuple[str, str], records: list[dict[str, Any]]) -> dict[str, Any]:
    identifier_type, token = key
    id_token = token if identifier_type == "cve" else f"{identifier_type.upper()}-{token.upper()}"
    case_group_id = f"GROUP-{implementation}-{id_token}"
    case_id = f"CASE-{implementation}-{id_token}-RC01"
    cves = ordered_unique([cve for record in records for cve in record["_cves"]])
    commits = ordered_unique([commit for record in records for commit in record["_commits"]])
    root_causes = _values(records, "root_cause")
    priorities_raw = _values(records, "priority_raw")
    priorities_normalized = ordered_unique(
        [PRIORITY_MAP.get(value, "unmapped") for value in priorities_raw]
    )
    title_values = _values(records, "title")
    aliases = ordered_unique(cves + commits)
    aliases = [alias for alias in aliases if alias.lower() != token.lower()]

    public_records = []
    references = []
    for record in records:
        references.extend(record["_references"])
        public_records.append({key: value for key, value in record.items() if not key.startswith("_")})

    identifier = {
        "type": identifier_type,
        "cve_id": token if identifier_type == "cve" else None,
        "commit_sha": token.lower() if identifier_type == "commit" else None,
        "aliases": aliases,
    }
    case = {
        "schema_version": SCHEMA_VERSION,
        "case_id": case_id,
        "case_group_id": case_group_id,
        "root_cause_index": 1,
        "source_implementation": implementation,
        "identifier": identifier,
        "facts": {
            "title": title_values[0] if len(title_values) == 1 else (" | ".join(title_values) or None),
            "target_implementations": _values(records, "target_implementation"),
            "protocol_versions_raw": _values(records, "protocol_version"),
            "priority": {"raw": priorities_raw, "normalized": priorities_normalized},
            "attack_surfaces_raw": _values(records, "attack_surface"),
            "root_cause_claims": root_causes,
            "affected_versions_raw": _values(records, "affected_versions"),
            "cvss_raw": _values(records, "cvss_raw"),
            "published_or_commit_dates": _values(records, "published_or_commit_date"),
        },
        "references": references,
        "source_records": public_records,
        "field_status": {
            "title": "known" if title_values else "unknown",
            "root_cause": (
                "unknown" if not root_causes else "known" if len(root_causes) == 1 else "needs_review"
            ),
            "identifier": "known" if identifier_type != "unresolved" else "needs_review",
        },
        "workflow": {
            "root_cause_split_status": "single_claim" if len(root_causes) <= 1 else "needs_review",
            "extraction_status": "not_started",
            "knfsd_mapping_status": "not_started",
        },
        "source_fingerprint_sha256": "",
    }
    case["source_fingerprint_sha256"] = sha256_bytes(canonical_json(public_records))
    return case


def ingest(repo: Path, input_dir: Path, output_dir: Path) -> dict[str, Any]:
    source_config = yaml.safe_load((repo / "config" / "sources.yaml").read_text(encoding="utf-8"))
    source_defs = source_config["sources"]
    expected = {item["filename"] for item in source_defs}
    missing = sorted(filename for filename in expected if not (input_dir / filename).is_file())
    if missing:
        raise FileNotFoundError(f"Missing configured workbooks: {missing}")

    case_schema = json.loads((repo / "schemas" / "case.schema.json").read_text(encoding="utf-8"))
    validator = Draft202012Validator(case_schema)
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    source_manifest = []

    for source in source_defs:
        implementation = source["implementation"]
        path = input_dir / source["filename"]
        workbook_sha256 = sha256_file(path)
        workbook = load_workbook(path, read_only=False, data_only=False)
        count = 0
        sheets = []
        for sheet in workbook.worksheets:
            headers = [str(sheet.cell(row=1, column=i).value).strip() for i in range(1, sheet.max_column + 1)]
            unknown = [header for header in headers if header not in HEADER_ALIASES]
            if unknown:
                raise ValueError(f"Unsupported headers in {path.name}/{sheet.title}: {unknown}")
            sheet_count = 0
            for row_index in range(2, sheet.max_row + 1):
                if not any(sheet.cell(row=row_index, column=i).value is not None for i in range(1, sheet.max_column + 1)):
                    continue
                record = _record(path, workbook_sha256, implementation, sheet, row_index, headers)
                kind, token = _group_key(record)
                grouped[(implementation, kind, token)].append(record)
                count += 1
                sheet_count += 1
            sheets.append({"name": sheet.title, "record_count": sheet_count, "headers": headers})
        workbook.close()
        source_manifest.append(
            {
                "implementation": implementation,
                "workbook": path.name,
                "sha256": workbook_sha256,
                "record_count": count,
                "sheets": sheets,
            }
        )

    cases = []
    for (implementation, kind, token), records in sorted(grouped.items()):
        case = _build_case(implementation, (kind, token), records)
        errors = sorted(validator.iter_errors(case), key=lambda error: list(error.path))
        if errors:
            details = "; ".join(f"{list(error.path)}: {error.message}" for error in errors[:5])
            raise ValueError(f"Generated invalid case {case['case_id']}: {details}")
        cases.append(case)

    case_dir = output_dir / "cases"
    case_dir.mkdir(parents=True, exist_ok=True)
    for stale in case_dir.glob("*.json"):
        stale.unlink()
    for case in cases:
        write_json(case_dir / f"{case['case_id']}.json", case)

    generated_dir = output_dir / "generated"
    generated_dir.mkdir(parents=True, exist_ok=True)
    jsonl_path = generated_dir / "cases.jsonl"
    jsonl_path.write_text(
        "".join(json.dumps(case, ensure_ascii=False, sort_keys=True) + "\n" for case in cases),
        encoding="utf-8",
    )
    identifier_counts = Counter(case["identifier"]["type"] for case in cases)
    status_counts = Counter(case["workflow"]["root_cause_split_status"] for case in cases)
    quality_report = {
        "schema_version": SCHEMA_VERSION,
        "case_count": len(cases),
        "identifier_counts": dict(sorted(identifier_counts.items())),
        "root_cause_split_status_counts": dict(sorted(status_counts.items())),
        "unresolved_case_ids": [
            case["case_id"] for case in cases if case["identifier"]["type"] == "unresolved"
        ],
        "root_cause_review_case_ids": [
            case["case_id"]
            for case in cases
            if case["workflow"]["root_cause_split_status"] == "needs_review"
        ],
    }
    write_json(output_dir / "quality-report.json", quality_report)
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source_count": len(source_manifest),
        "source_record_count": sum(item["record_count"] for item in source_manifest),
        "case_count": len(cases),
        "sources": source_manifest,
        "jsonl": {"path": "generated/cases.jsonl", "sha256": sha256_file(jsonl_path)},
        "semantics": {
            "case_unit": "one source implementation plus one primary CVE or commit, initially RC01",
            "facts_only": True,
            "root_cause_split": "multiple claims are flagged for independent LLM review; inferred splits are stored outside canonical facts",
        },
    }
    write_json(output_dir / "manifest.json", manifest)
    return manifest


def main() -> None:
    args = _args()
    repo = args.repo.resolve()
    input_dir = (args.input_dir or repo / "data" / "raw").resolve()
    output_dir = (args.output_dir or repo / "data").resolve()
    result = ingest(repo, input_dir, output_dir)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
